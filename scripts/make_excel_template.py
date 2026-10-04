"""IT ARENA uchun Excel savollar shabloni va Informatika 9-sinf bankining Excel nusxasini yaratadi.

Ishga tushirish:  python scripts/make_excel_template.py
Natija:
  public/templates/IT_ARENA_savollar_shabloni.xlsx       — boʻsh shablon + namunalar + yoʻriqnoma
  public/templates/Informatika_9-sinf_savollar_banki.xlsx — tayyor 117 ta savol (tahrirlab qayta yuklash mumkin)
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "public" / "templates"

COLUMNS = [
    ("Fan", 14, "Fan nomi: Informatika, Matematika, Ona tili, Fizika ... Boʻsh qolsa, import oynasida tanlangan fan olinadi."),
    ("Sinf", 7, "Sinf raqami: 5, 6, 7, 8, 9, 10 yoki 11. Boʻsh qolsa — barcha sinflar uchun."),
    ("Mavzu", 30, "Darsdagi mavzu nomi. Xona yaratishda mavzular boʻyicha savol tanlash mumkin."),
    ("Savol turi", 22, "Roʻyxatdan tanlang: Test, Toʻgʻri/Notoʻgʻri, Rasmli, Qisqa javob, Mantiqiy."),
    ("Savol matni", 60, "Savolning oʻzi (3–600 belgi)."),
    ("A variant", 26, "Test, Rasmli va Mantiqiy savollar uchun variantlar. Kamida A va B toʻldirilsin."),
    ("B variant", 26, None),
    ("C variant", 26, None),
    ("D variant", 26, None),
    ("Toʻgʻri javob", 16, "Test/Rasmli/Mantiqiy: A, B, C yoki D.\nToʻgʻri/Notoʻgʻri: Toʻgʻri yoki Notoʻgʻri.\nQisqa javob: javob(lar), bir nechta boʻlsa ; bilan ajrating (masalan: 6; olti)."),
    ("Izoh", 44, "Javob ochilganda oʻquvchilarga koʻrsatiladigan qisqa tushuntirish."),
    ("Maslahat", 30, "HINT SCAN qobiliyati ishlatilganda jamoaga koʻrsatiladi (ixtiyoriy)."),
    ("Rasm havolasi", 30, "Rasmli savol uchun: https:// bilan boshlanuvchi havola. Rasmni kompyuterdan yuklash uchun importdan keyin savolni tahrirlab, “Rasm yuklash” tugmasidan foydalaning."),
    ("Qiyinlik", 10, "oson, oʻrta yoki qiyin."),
    ("Raund", 8, "Tavsiya etilgan raund 1–5 (ixtiyoriy):\n1 — Bilimlar sinovi\n2 — Rasmli ov\n3 — Tezlik jangi\n4 — Tizimni buz (mantiq)\n5 — Galaktik boss (qiyin)"),
    ("Ball", 8, "Ixtiyoriy. Boʻsh qolsa raund boʻyicha: 100/150/200/250/300."),
    ("Vaqt (soniya)", 12, "Ixtiyoriy, 5–300 soniya. Boʻsh qolsa raund boʻyicha."),
]

TYPE_LIST = ["Test", "Toʻgʻri/Notoʻgʻri", "Rasmli", "Qisqa javob", "Mantiqiy"]
DIFF_LIST = ["oson", "oʻrta", "qiyin"]
TYPE_LABEL = {
    "single_choice": "Test",
    "true_false": "Toʻgʻri/Notoʻgʻri",
    "image_identification": "Rasmli",
    "short_answer": "Qisqa javob",
    "logical_puzzle": "Mantiqiy",
}
DIFF_LABEL = {"easy": "oson", "medium": "oʻrta", "hard": "qiyin"}

EXAMPLES = [
    ["Informatika", 9, "Kiritish qurilmalari", "Test", "Quyidagilardan qaysi biri kiritish qurilmasi?",
     "Monitor", "Mikrofon", "Printer", "Karnay", "B",
     "Mikrofon ovozni kompyuterga kiritadi.", "U ovozni “tinglaydi”.", "", "oson", 1, "", ""],
    ["Matematika", 7, "Oddiy kasrlar", "Test", "1/2 + 1/4 nechaga teng?",
     "2/6", "3/4", "1/6", "2/4", "B",
     "1/2 = 2/4, demak 2/4 + 1/4 = 3/4.", "Umumiy maxrajga keltiring.", "", "oʻrta", 1, "", ""],
    ["Ona tili", 5, "Unli tovushlar", "Toʻgʻri/Notoʻgʻri", "Oʻzbek tilida 6 ta unli tovush bor.",
     "", "", "", "", "Toʻgʻri",
     "Unli tovushlar: a, o, u, e, i, oʻ.", "", "", "oson", 3, "", ""],
    ["Fizika", 8, "Kuch va harakat", "Qisqa javob", "Kuchning SI tizimidagi birligi qanday ataladi?",
     "", "", "", "", "Nyuton; N",
     "Kuch Nyuton (N) da oʻlchanadi.", "Olim familiyasi bilan atalgan.", "", "oʻrta", 4, "", ""],
    ["Informatika", 9, "Kiritish qurilmalari", "Rasmli", "Rasmda qanday qurilma tasvirlangan?",
     "Grafik planshet", "Klaviatura", "Skaner", "Sensorli panel", "B",
     "Klaviatura — matn kiritish qurilmasi.", "", "/assets/devices/keyboard.webp", "oson", 2, "", ""],
    ["Ingliz tili", 10, "Present Perfect", "Mantiqiy", "Choose the correct form: She ___ already ___ her homework.",
     "has / finished", "have / finished", "is / finishing", "did / finish", "A",
     "He/she/it bilan Present Perfect: has + V3.", "Ega 3-shaxs birlikda.", "", "qiyin", 5, 300, 45],
]

DARK = PatternFill("solid", fgColor="0E1B2D")
EXAMPLE_FILL = PatternFill("solid", fgColor="EAF7FF")
HEADER_FONT = Font(bold=True, color="3EE7FF", size=11)
THIN = Side(style="thin", color="B7C7DA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_questions_sheet(ws, rows, mark_examples=False, total_rows=300):
    ws.title = "Savollar"
    for idx, (title, width, note) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = DARK
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        if note:
            cell.comment = Comment(note, "IT ARENA", width=320, height=140)
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:Q{max(2, len(rows) + 1)}"

    for r, values in enumerate(rows, start=2):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value if value != "" else None)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if mark_examples:
                cell.fill = EXAMPLE_FILL

    last = max(total_rows, len(rows) + 50)
    dv_type = DataValidation(type="list", formula1='"' + ",".join(TYPE_LIST) + '"', allow_blank=True,
                             showErrorMessage=True, errorTitle="Savol turi", error="Roʻyxatdan tanlang")
    dv_diff = DataValidation(type="list", formula1='"' + ",".join(DIFF_LIST) + '"', allow_blank=True,
                             showErrorMessage=True, errorTitle="Qiyinlik", error="oson, oʻrta yoki qiyin")
    dv_round = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    dv_grade = DataValidation(type="whole", operator="between", formula1="1", formula2="11", allow_blank=True,
                              showErrorMessage=True, errorTitle="Sinf", error="1 dan 11 gacha son kiriting")
    dv_points = DataValidation(type="whole", operator="between", formula1="0", formula2="2000", allow_blank=True)
    dv_time = DataValidation(type="whole", operator="between", formula1="5", formula2="300", allow_blank=True)
    for dv, col in ((dv_type, "D"), (dv_diff, "N"), (dv_round, "O"), (dv_grade, "B"), (dv_points, "P"), (dv_time, "Q")):
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")


def build_instructions(wb):
    ws = wb.create_sheet("Yoʻriqnoma")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    title = ws.cell(row=1, column=1, value="IT ARENA — Galaktik jang: savollar shabloni")
    title.font = Font(bold=True, size=16, color="2583FF")
    ws.cell(row=2, column=1, value="Bu shablon istalgan fan va 5–11-sinflar uchun ishlaydi. "
                                    "“Savollar” varagʻini toʻldiring va oʻqituvchi panelida yuklang.").font = Font(italic=True)
    steps = [
        ("1-qadam", "“Savollar” varagʻidagi rangli namuna qatorlarini oʻchiring (yoki qoldiring — import oynasida belgini olib tashlash mumkin)."),
        ("2-qadam", "Har bir savol — bitta qator. Fan, Sinf va Mavzuni yozing: xona yaratishda shular boʻyicha savollar tanlanadi."),
        ("3-qadam", "“Savol turi” va “Qiyinlik” ustunlarida ochiluvchi roʻyxatdan tanlang."),
        ("4-qadam", "Faylni .xlsx formatida saqlang."),
        ("5-qadam", "Oʻqituvchi paneli → Savollar banki → “Excel yuklash” → faylni tanlang → tekshiruv natijasini koʻring → “Saqlash”."),
    ]
    row = 4
    for k, v in steps:
        ws.cell(row=row, column=1, value=k).font = Font(bold=True)
        ws.cell(row=row, column=2, value=v).alignment = Alignment(wrap_text=True)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Ustunlar").font = Font(bold=True, size=13, color="2583FF")
    row += 1
    for title_text, _, note in COLUMNS:
        ws.cell(row=row, column=1, value=title_text).font = Font(bold=True)
        cell = ws.cell(row=row, column=2, value=note or "Test variantlari (ixtiyoriy C va D).")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Savol turlari").font = Font(bold=True, size=13, color="2583FF")
    row += 1
    types = [
        ("Test", "2–4 variant, bitta toʻgʻri javob (A/B/C/D)."),
        ("Toʻgʻri/Notoʻgʻri", "Variantlar avtomatik qoʻyiladi. Javob: Toʻgʻri yoki Notoʻgʻri."),
        ("Rasmli", "Test kabi, lekin katta ekranda rasm koʻrsatiladi. “Rasm havolasi” majburiy."),
        ("Qisqa javob", "Oʻquvchi javobni yozadi (raqamli qulf ham shu tur). Bir nechta toʻgʻri variant ; bilan."),
        ("Mantiqiy", "Mantiqiy masala yoki ketma-ketlik, variantli (A/B/C/D)."),
    ]
    for k, v in types:
        ws.cell(row=row, column=1, value=k).font = Font(bold=True)
        ws.cell(row=row, column=2, value=v).alignment = Alignment(wrap_text=True)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Maslahatlar").font = Font(bold=True, size=13, color="2583FF")
    row += 1
    tips = [
        "Toʻgʻri javoblar faqat serverda saqlanadi va oʻquvchilarga javob vaqti tugaguncha yuborilmaydi.",
        "Toʻgʻri javob har doim A boʻlib qolmasin — variantlarni aralashtiring.",
        "Rasmlarni kompyuterdan yuklash: importdan keyin savolni tahrirlang → “Rasm yuklash”.",
        "Bir xil savolni ikki marta yuklamaslik uchun import oynasi takroriy savollarni belgilaydi.",
        "Eski .xls emas, .xlsx formatidan foydalaning. CSV (UTF-8) ham qabul qilinadi.",
    ]
    for t in tips:
        ws.cell(row=row, column=1, value="•")
        ws.cell(row=row, column=2, value=t).alignment = Alignment(wrap_text=True)
        row += 1


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    build_questions_sheet(wb.active, EXAMPLES, mark_examples=True)
    build_instructions(wb)
    template_path = OUT_DIR / "IT_ARENA_savollar_shabloni.xlsx"
    wb.save(template_path)
    print(f"✔ {template_path.relative_to(ROOT)}")

    bank = json.loads((ROOT / "supabase" / "seed" / "question_bank.json").read_text(encoding="utf8"))
    rows = []
    letters = "ABCDEF"
    for q in bank:
        t = q["question_type"]
        opts = [] if t == "true_false" else q["options"]
        if t == "short_answer":
            correct = "; ".join(q["correct_answer"])
        elif t == "true_false":
            correct = "Toʻgʻri" if q["correct_answer"] == 0 else "Notoʻgʻri"
        else:
            correct = letters[q["correct_answer"]]
        rows.append([
            q["subject"], q["grade"], q["category"], TYPE_LABEL[t], q["question_text"],
            *(opts + ["", "", "", ""])[:4], correct, q["explanation"], q["hint"], q["image_url"] or "",
            DIFF_LABEL[q["difficulty"]], q["recommended_round"], q["default_points"], q["default_time_limit"],
        ])
    wb2 = Workbook()
    build_questions_sheet(wb2.active, rows, mark_examples=False)
    build_instructions(wb2)
    bank_path = OUT_DIR / "Informatika_9-sinf_savollar_banki.xlsx"
    wb2.save(bank_path)
    print(f"✔ {bank_path.relative_to(ROOT)} ({len(rows)} ta savol)")


if __name__ == "__main__":
    main()
